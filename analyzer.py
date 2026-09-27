"""OPTCG deck analyzer.

Pulls tournament decklists from the Limitless API and card stats from
OPTCG API into a local SQLite database, then writes a per-deck breakdown
to an Excel workbook, grouped by set.
"""
import argparse
import re
import sqlite3
import sys
import time
from collections import Counter, defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).parent
DB_PATH = ROOT / "data" / "optcg.db"
LIMITLESS = "https://play.limitlesstcg.com/api"
OPTCGAPI = "https://optcgapi.com/api"
CARD_CACHE_MAX_AGE = timedelta(days=7)
REQUEST_GAP = 1.0  # seconds between standings requests, to stay under Limitless's rate limit
SET_MIN_DECKS = 3  # a set has "started" at the first tournament where this many decks play its cards
MAX_COST = 10

# Timing tags that end an [On Play] / [Main] section of card text.
_TAG = r"\[(?:On Play|When Attacking|Trigger|Main|Activate: ?Main|Counter|On K\.O\.|End of Your Turn|On Block|Your Turn|Opponent's Turn|On Your Opponent's Attack)\]"
_ON_PLAY = re.compile(r"\[On Play\](.*?)(?=" + _TAG + r"|$)", re.S)
_MAIN = re.compile(r"\[Main\](.*?)(?=" + _TAG + r"|$)", re.S)
_LOOK = re.compile(r"(?:look at|search) .*?cards? from the top of your deck|search your deck", re.I)
_TAKE = re.compile(r"reveal up to|add (?:it|them|up to \d+[^.]*?) to your hand", re.I)
# OPTCG API decorates names: "Nami (062)", "Sabo (004) (Alternate Art)", "Dracule Mihawk - OP14-020".
_NAME_SUFFIX = re.compile(r"(?:\s*\([^()]*\)|\s*-\s*[A-Z0-9]+-\d+)+\s*$")
_FORMAT_TAG = re.compile(r"\[([A-Z0-9]+)\]")
# Tournament formats that are not standard constructed; kept apart from the set's normal events.
SIDE_FORMATS = {"EGB"}

SCHEMA = """
CREATE TABLE IF NOT EXISTS meta (key TEXT PRIMARY KEY, value TEXT);
CREATE TABLE IF NOT EXISTS card (
    id TEXT PRIMARY KEY, name TEXT, type TEXT, color TEXT,
    cost INTEGER, power INTEGER, counter INTEGER NOT NULL DEFAULT 0, text TEXT,
    on_play_searcher INTEGER NOT NULL, event_searcher INTEGER NOT NULL, counter_event INTEGER NOT NULL
);
CREATE TABLE IF NOT EXISTS tournament (
    id TEXT PRIMARY KEY, name TEXT, date TEXT, format TEXT, players INTEGER,
    complete INTEGER NOT NULL DEFAULT 0
);
CREATE TABLE IF NOT EXISTS entry (
    id INTEGER PRIMARY KEY, tournament_id TEXT NOT NULL REFERENCES tournament(id),
    player TEXT, placing INTEGER, wins INTEGER, losses INTEGER, ties INTEGER,
    leader_id TEXT, leader_name TEXT
);
CREATE TABLE IF NOT EXISTS deck_card (
    entry_id INTEGER NOT NULL REFERENCES entry(id), card_id TEXT NOT NULL, count INTEGER NOT NULL,
    PRIMARY KEY (entry_id, card_id)
);
-- Set eras detected from the decklists (see update_set_eras). end_date is NULL for the current set.
CREATE TABLE IF NOT EXISTS set_era (
    code TEXT PRIMARY KEY, start_date TEXT NOT NULL, end_date TEXT, partial INTEGER NOT NULL
);
CREATE INDEX IF NOT EXISTS entry_tournament ON entry(tournament_id);
CREATE INDEX IF NOT EXISTS entry_leader ON entry(leader_id);

-- One row per deck with the headline numbers, for ad-hoc and trend queries.
-- Same definitions as analyze() below; cards missing from `card` are not counted.
DROP VIEW IF EXISTS deck_summary;
CREATE VIEW deck_summary AS
SELECT e.id AS entry_id,
       (SELECT code FROM set_era s WHERE t.date >= s.start_date AND (s.end_date IS NULL OR t.date < s.end_date)) AS set_code,
       t.date, t.format, t.name AS tournament, e.player, e.placing,
       e.wins, e.losses, e.leader_id, e.leader_name,
       SUM(CASE WHEN c.type = 'Character' THEN dc.count ELSE 0 END) AS characters,
       SUM(CASE WHEN c.type = 'Event' THEN dc.count ELSE 0 END) AS events,
       SUM(CASE WHEN c.type = 'Stage' THEN dc.count ELSE 0 END) AS stages,
       SUM(CASE WHEN c.counter = 0 THEN dc.count ELSE 0 END) AS bricks,
       SUM(CASE WHEN c.counter = 0 AND c.counter_event THEN dc.count ELSE 0 END) AS bricks_counter_events,
       SUM(CASE WHEN c.counter > 0 THEN dc.count ELSE 0 END) AS with_counter,
       SUM(CASE WHEN c.counter = 1000 THEN dc.count ELSE 0 END) AS counter_1k,
       SUM(CASE WHEN c.counter = 2000 THEN dc.count ELSE 0 END) AS counter_2k,
       SUM(c.counter * dc.count) AS counter_total,
       SUM(CASE WHEN c.on_play_searcher THEN dc.count ELSE 0 END) AS searchers,
       SUM(CASE WHEN c.event_searcher THEN dc.count ELSE 0 END) AS event_searchers,
       ROUND(1.0 * SUM(MIN(c.cost, 10) * dc.count) / SUM(CASE WHEN c.cost IS NOT NULL THEN dc.count END), 2) AS avg_cost
FROM entry e
JOIN tournament t ON t.id = e.tournament_id
JOIN deck_card dc ON dc.entry_id = e.id
JOIN card c ON c.id = dc.card_id AND c.type != 'Leader'
GROUP BY e.id;
"""


def connect():
    DB_PATH.parent.mkdir(exist_ok=True)
    db = sqlite3.connect(DB_PATH)
    db.row_factory = sqlite3.Row
    db.executescript(SCHEMA)
    return db


def _get(url, **params):
    for attempt in range(6):
        r = requests.get(url, params=params, timeout=30)
        if r.status_code == 429:
            wait = int(r.headers.get("Retry-After") or 0) or 15 * (attempt + 1)
            print(f"  Rate limited by {url.split('/')[2]}, waiting {wait}s...", file=sys.stderr)
            time.sleep(wait)
            continue
        r.raise_for_status()
        return r.json()
    sys.exit("Still rate limited after several retries. Wait a few minutes and run again; progress so far is saved.")


def _date(s):
    return datetime.fromisoformat(s.replace("Z", "+00:00"))


# ---------------------------------------------------------------- cards

def _num(v):
    try:
        return int(v)
    except (TypeError, ValueError):
        return None


def _is_searcher(text, pattern):
    return any(_LOOK.search(seg) and _TAKE.search(seg) for seg in pattern.findall(text or ""))


def load_cards(db, refresh=False):
    """Return {card_id: card dict}, re-downloading from OPTCG API when stale."""
    row = db.execute("SELECT value FROM meta WHERE key = 'cards_updated'").fetchone()
    stale = not row or datetime.now() - datetime.fromisoformat(row["value"]) > CARD_CACHE_MAX_AGE
    if stale or refresh:
        print("Downloading card database from OPTCG API...", file=sys.stderr)
        raw = []
        for endpoint in ("allSetCards", "allSTCards", "allPromos"):
            raw += _get(f"{OPTCGAPI}/{endpoint}/")
        seen = {}
        for c in raw:
            if c["card_set_id"] in seen:  # parallel/alt-art reprint, same stats
                continue
            text = c.get("card_text") or ""
            seen[c["card_set_id"]] = (
                c["card_set_id"], _NAME_SUFFIX.sub("", c["card_name"]), c["card_type"], c.get("card_color"),
                _num(c.get("card_cost")), _num(c.get("card_power")), _num(c.get("counter_amount")) or 0, text,
                c["card_type"] == "Character" and _is_searcher(text, _ON_PLAY),
                c["card_type"] == "Event" and _is_searcher(text, _MAIN),
                c["card_type"] == "Event" and "[Counter]" in text,
            )
        with db:
            db.executemany("INSERT OR REPLACE INTO card VALUES (?,?,?,?,?,?,?,?,?,?,?)", seen.values())
            db.execute("INSERT OR REPLACE INTO meta VALUES ('cards_updated', ?)", (datetime.now().isoformat(),))
    return {r["id"]: dict(r) for r in db.execute("SELECT * FROM card")}


# ---------------------------------------------------------------- sync

def fetch_tournaments(event, fmt, cutoff, min_players):
    found, page = [], 1
    while True:
        batch = _get(f"{LIMITLESS}/tournaments", game="OP", limit=100, page=page)
        if not batch:
            break
        for t in batch:
            name = t["name"].lower()
            if event and not any(e.lower() in name for e in event):
                continue
            if fmt and fmt.lower() not in name:
                continue
            if (t.get("players") or 0) < min_players:
                continue
            if _date(t["date"]) >= cutoff:
                found.append(t)
        if _date(batch[-1]["date"]) < cutoff:
            break
        page += 1
    return found


def _card_id(c):
    return f"{c['set']}-{c['number']}"


def store_tournament(db, t, standings):
    # Events less than a day old may still be running; mark them so the next sync refetches.
    complete = datetime.now(timezone.utc) - _date(t["date"]) > timedelta(days=1)
    m = _FORMAT_TAG.search(t["name"])
    with db:
        db.execute("DELETE FROM deck_card WHERE entry_id IN (SELECT id FROM entry WHERE tournament_id = ?)", (t["id"],))
        db.execute("DELETE FROM entry WHERE tournament_id = ?", (t["id"],))
        db.execute("INSERT OR REPLACE INTO tournament VALUES (?,?,?,?,?,?)",
                   (t["id"], t["name"], t["date"], m.group(1) if m else None, t.get("players"), complete))
        for p in standings:
            deck = p.get("decklist")
            if not deck or not deck.get("leader"):
                continue
            r = p.get("record") or {}
            cur = db.execute(
                "INSERT INTO entry (tournament_id, player, placing, wins, losses, ties, leader_id, leader_name) "
                "VALUES (?,?,?,?,?,?,?,?)",
                (t["id"], p["name"], p.get("placing"), r.get("wins", 0), r.get("losses", 0), r.get("ties", 0),
                 _card_id(deck["leader"]), deck["leader"]["name"]))
            counts = Counter()
            for group, items in deck.items():
                if group != "leader":
                    for c in items:
                        counts[_card_id(c)] += c["count"]
            db.executemany("INSERT INTO deck_card VALUES (?,?,?)",
                           [(cur.lastrowid, cid, n) for cid, n in counts.items()])


def sync_cutoff(db, args):
    """--days if given; otherwise a week before the newest matching event we have; otherwise all history."""
    if args.days:
        return datetime.now(timezone.utc) - timedelta(days=args.days)
    where = " OR ".join("name LIKE ?" for _ in args.event)
    row = db.execute(f"SELECT MAX(date) AS d FROM tournament WHERE {where}", [f"%{e}%" for e in args.event]).fetchone()
    if row["d"]:
        return _date(row["d"]) - timedelta(days=7)
    print("First sync: downloading full history, this takes a few minutes...", file=sys.stderr)
    return datetime(2000, 1, 1, tzinfo=timezone.utc)


def sync(db, args):
    tournaments = fetch_tournaments(args.event, args.format, sync_cutoff(db, args), args.min_players)
    done = {r["id"] for r in db.execute("SELECT id FROM tournament WHERE complete = 1")}
    # Oldest first: if a long sync is interrupted, the next run resumes from the newest stored event.
    new = sorted((t for t in tournaments if t["id"] not in done), key=lambda t: t["date"])
    for i, t in enumerate(new, 1):
        print(f"  [{i}/{len(new)}] {t['date'][:10]} {t['name']}", file=sys.stderr)
        store_tournament(db, t, _get(f"{LIMITLESS}/tournaments/{t['id']}/standings"))
        if i < len(new):
            time.sleep(REQUEST_GAP)
    print(f"Synced: {len(new)} new/updated, {len(tournaments) - len(new)} already in database.", file=sys.stderr)
    update_set_eras(db)


# ---------------------------------------------------------------- sets

def update_set_eras(db):
    """Work out when each main set (OPxx) started, from the first tournament where it was played.

    A set starts at the first standard tournament where at least SET_MIN_DECKS decks play one
    of its cards, and ends when the next set starts. Sets already being played at the oldest
    tournament in the database started earlier than our data, so only the newest of those is
    kept, marked partial.
    """
    first = db.execute(f"""
        SELECT set_code, MIN(date) AS start_date FROM (
            SELECT substr(dc.card_id, 1, 4) AS set_code, t.id, t.date
            FROM deck_card dc JOIN entry e ON e.id = dc.entry_id JOIN tournament t ON t.id = e.tournament_id
            WHERE dc.card_id GLOB 'OP[0-9][0-9]-*' AND COALESCE(t.format, '') NOT IN ({','.join('?' * len(SIDE_FORMATS))})
            GROUP BY set_code, t.id HAVING COUNT(DISTINCT e.id) >= {SET_MIN_DECKS})
        GROUP BY set_code ORDER BY start_date, set_code""", list(SIDE_FORMATS)).fetchall()
    if not first:
        return
    oldest = db.execute("SELECT MIN(date) AS d FROM tournament").fetchone()["d"]
    predating = [r for r in first if r["start_date"] <= oldest]
    eras = ([(predating[-1]["set_code"], oldest, True)] if predating else []) + \
           [(r["set_code"], r["start_date"], False) for r in first if r["start_date"] > oldest]
    with db:
        db.execute("DELETE FROM set_era")
        for i, (code, start, partial) in enumerate(eras):
            end = eras[i + 1][1] if i + 1 < len(eras) else None
            db.execute("INSERT INTO set_era VALUES (?,?,?,?)", (code, start, end, partial))


def load_set_eras(db):
    return [dict(r) for r in db.execute("SELECT * FROM set_era ORDER BY start_date")]


def set_label(eras, date, fmt):
    code = next((e["code"] for e in eras if date >= e["start_date"] and (not e["end_date"] or date < e["end_date"])), "?")
    return f"{code} · {fmt}" if fmt in SIDE_FORMATS else code


# ---------------------------------------------------------------- analysis

def load_entries(db, args):
    eras = load_set_eras(db)
    where, params = ["t.players >= ?"], [args.min_players]
    if args.days:
        where.append("t.date >= ?")
        params.append((datetime.now(timezone.utc) - timedelta(days=args.days)).strftime("%Y-%m-%d"))
    if args.event:
        where.append("(" + " OR ".join("t.name LIKE ?" for _ in args.event) + ")")
        params += [f"%{e}%" for e in args.event]
    if args.format:
        where.append("t.name LIKE ?")
        params.append(f"%{args.format}%")
    if args.top:
        where.append("e.placing <= ?")
        params.append(args.top)
    if args.leader:
        where.append("(e.leader_name LIKE ? OR e.leader_id LIKE ?)")
        params += [f"%{args.leader}%"] * 2
    rows = db.execute(
        "SELECT e.*, t.name AS tournament, t.date AS iso_date, t.format FROM entry e "
        "JOIN tournament t ON t.id = e.tournament_id WHERE " + " AND ".join(where) +
        " ORDER BY t.date DESC, e.placing", params).fetchall()

    wanted_set = None
    if args.set and args.set.lower() != "all":
        wanted_set = eras[-1]["code"] if args.set.lower() == "current" and eras else args.set.upper()
    entries, ids = [], []
    for r in rows:
        label = set_label(eras, r["iso_date"], r["format"])
        if wanted_set and not label.startswith(wanted_set):
            continue
        ids.append(r["id"])
        entries.append({
            "id": r["id"], "tournament_id": r["tournament_id"], "set": label, "leader_name": f"{r['leader_name']} ({r['leader_id']})",
            "player": r["player"], "placing": r["placing"], "record": f"{r['wins']}-{r['losses']}-{r['ties']}",
            "tournament": r["tournament"], "date": r["iso_date"][:10],
        })
    counts = defaultdict(Counter)
    for chunk in range(0, len(ids), 500):
        part = ids[chunk:chunk + 500]
        for dc in db.execute(f"SELECT * FROM deck_card WHERE entry_id IN ({','.join('?' * len(part))})", part):
            counts[dc["entry_id"]][dc["card_id"]] = dc["count"]
    for e in entries:
        e["counts"] = counts[e["id"]]
    return entries, eras


def analyze(counts, cards):
    s = Counter()
    curve = Counter()
    unknown = []
    for cid, n in counts.items():
        c = cards.get(cid)
        if not c:
            unknown.append(cid)
            continue
        if c["type"] == "Leader":
            continue
        s["total"] += n
        s[c["type"].lower() + "s"] += n
        if c["counter"] > 0:
            s["with_counter"] += n
            s[f"counter_{c['counter'] // 1000}k"] += n
            s["counter_total"] += c["counter"] * n
        else:
            s["bricks"] += n
            if c["counter_event"]:
                s["bricks_counter_events"] += n
        if c["on_play_searcher"]:
            s["searchers"] += n
        if c["event_searcher"]:
            s["event_searchers"] += n
        if c["cost"] is not None:
            curve[min(c["cost"], MAX_COST)] += n
    s["avg_cost"] = round(sum(k * v for k, v in curve.items()) / max(sum(curve.values()), 1), 2)
    return s, curve, unknown


# ---------------------------------------------------------------- excel

DECK_COLS = [
    ("Characters", "characters"), ("Events", "events"), ("Stages", "stages"),
    ("Bricks (no counter)", "bricks"), ("  of which [Counter] events", "bricks_counter_events"),
    ("Cards w/ counter", "with_counter"), ("+1k counters", "counter_1k"), ("+2k counters", "counter_2k"),
    ("Total counter value", "counter_total"),
    ("On Play searchers", "searchers"), ("Event searchers", "event_searchers"),
    ("Avg cost", "avg_cost"),
]
COST_COLS = [f"{i}{'+' if i == MAX_COST else ''} cost" for i in range(MAX_COST + 1)]
TYPE_ORDER = {"Character": 0, "Event": 1, "Stage": 2}
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _sheet(wb, title, header, rows, widths=None, freeze="B2"):
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append(r)
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, max(10, min(len(str(h)) + 2, 28)))
    return ws


def _deck_row(s, curve):
    return [s.get(k, 0) for _, k in DECK_COLS] + [curve.get(i, 0) for i in range(MAX_COST + 1)]


def _yes(flag):
    return "Yes" if flag else ""


def write_workbook(entries, cards, out, eras=(), selection=""):
    wb = Workbook()
    wb.remove(wb.active)
    for e in entries:
        e["stats"], e["curve"], e["unknown"] = analyze(e["counts"], cards)

    # Sets: which sets and date ranges the report covers
    if eras:
        per_set = defaultdict(lambda: [set(), 0])
        for e in entries:
            per_set[e["set"]][0].add(e["tournament_id"])
            per_set[e["set"]][1] += 1
        rows = []
        for era in eras:
            for label in sorted(k for k in per_set if k.split(" ")[0] == era["code"]):
                rows.append([label, era["start_date"][:10],
                             "(current)" if not era["end_date"] else (_date(era["end_date"]) - timedelta(days=1)).strftime("%Y-%m-%d"),
                             len(per_set[label][0]), per_set[label][1],
                             "Set started before our data; start date is our oldest tournament" if era["partial"] else ""])
        ws = _sheet(wb, "Sets", ["Set", "From", "To", "Tournaments", "Decks in report", "Note"], rows,
                    {"Note": 60}, freeze="A2")
        ws.append([])
        ws.append([f"Selection: {selection}"])
        ws.append(["A set starts at the first standard tournament where 3+ decks play its cards, and ends when the next set starts. "
                   "EGB events are a separate format and listed on their own."])

    # Decks: one row per decklist
    header = ["Set", "Leader", "Player", "Placing", "Record", "Tournament", "Date"] + [c for c, _ in DECK_COLS] + COST_COLS + ["Unknown cards"]
    rows = [[e.get("set", ""), e["leader_name"], e["player"], e["placing"], e["record"], e["tournament"], e["date"]]
            + _deck_row(e["stats"], e["curve"]) + [", ".join(e["unknown"])] for e in entries]
    _sheet(wb, "Decks", header, rows, {"Leader": 26, "Tournament": 34, "Player": 16}, freeze="C2")

    # Decklists: every card in every deck
    header = ["Set", "Date", "Tournament", "Placing", "Player", "Leader", "Card ID", "Card name", "Copies",
              "Type", "Cost", "Counter", "Power", "Color", "Brick (no counter)", "Searcher", "[Counter] event"]
    rows = []
    for e in entries:
        def order(item):
            c = cards.get(item[0], {})
            return TYPE_ORDER.get(c.get("type"), 9), c.get("cost") if c.get("cost") is not None else 99, item[0]
        for cid, n in sorted(e["counts"].items(), key=order):
            c = cards.get(cid)
            if c and c["type"] == "Leader":
                continue
            if not c:
                rows.append([e.get("set", ""), e["date"], e["tournament"], e["placing"], e["player"], e["leader_name"],
                             cid, "(unknown card)", n] + [""] * 8)
                continue
            rows.append([e.get("set", ""), e["date"], e["tournament"], e["placing"], e["player"], e["leader_name"],
                         cid, c["name"], n, c["type"], c["cost"], c["counter"], c["power"], c["color"],
                         _yes(c["counter"] == 0), _yes(c["on_play_searcher"] or c["event_searcher"]), _yes(c["counter_event"])])
    _sheet(wb, "Decklists", header, rows, {"Tournament": 34, "Leader": 26, "Player": 16, "Card name": 30}, freeze="H2")

    # Leaders: averages per set + leader
    by_leader = defaultdict(list)
    for e in entries:
        by_leader[(e.get("set", ""), e["leader_name"])].append(e)
    header = ["Set", "Leader", "Decks", "Best placing"] + [c for c, _ in DECK_COLS] + COST_COLS
    rows = []
    for (set_code, leader), es in sorted(by_leader.items(), key=lambda kv: (kv[0][0], -len(kv[1]))):
        n = len(es)
        def avg(vals):
            return round(sum(vals) / n, 1)
        placings = [e["placing"] for e in es if e["placing"]]
        rows.append([set_code, leader, n, min(placings) if placings else None]
                    + [avg(e["stats"].get(k, 0) for e in es) for _, k in DECK_COLS]
                    + [avg(e["curve"].get(i, 0) for e in es) for i in range(MAX_COST + 1)])
    _sheet(wb, "Leaders", header, rows, {"Leader": 26}, freeze="C2")

    # Card usage per set + leader: core vs flex
    header = ["Set", "Leader", "Card ID", "Card name", "Type", "Cost", "Counter", "Searcher", "% of decks", "Avg copies (when played)"]
    rows = []
    for (set_code, leader), es in sorted(by_leader.items(), key=lambda kv: (kv[0][0], -len(kv[1]))):
        usage = defaultdict(list)
        for e in es:
            for cid, n in e["counts"].items():
                usage[cid].append(n)
        for cid, ns in sorted(usage.items(), key=lambda kv: (-len(kv[1]), -sum(kv[1]))):
            c = cards.get(cid, {})
            if c.get("type") == "Leader":
                continue
            rows.append([set_code, leader, cid, c.get("name", "(unknown card)"), c.get("type"), c.get("cost"), c.get("counter"),
                         _yes(c.get("on_play_searcher") or c.get("event_searcher")),
                         round(100 * len(ns) / len(es)), round(sum(ns) / len(ns), 2)])
    _sheet(wb, "Card Usage", header, rows, {"Leader": 26, "Card name": 30}, freeze="C2")

    # Definitions so the numbers are unambiguous
    defs = [
        ("Set", "Which set was newest when the tournament was played. EGB events are a separate format, labelled e.g. 'OP17 · EGB'."),
        ("Bricks (no counter)", "Non-leader cards with no counter value (0 or blank). Includes all events and stages."),
        ("  of which [Counter] events", "Bricks that are events with a [Counter] effect, i.e. still usable on defense."),
        ("Cards w/ counter", "Non-leader cards with a +1000 or +2000 counter value."),
        ("On Play searchers", "Characters whose [On Play] looks at/searches the deck AND reveals or adds a card to hand. Pure 'look and reorder' effects are excluded."),
        ("Event searchers", "Events whose [Main] effect searches the deck the same way. Counted separately."),
        ("N cost", "Number of non-leader cards at that DON!! cost. The last column groups 10+."),
        ("Unknown cards", "Card IDs not found in OPTCG API (usually brand-new cards). They are left out of all counts."),
    ]
    _sheet(wb, "Definitions", ["Column", "Meaning"], defs, {"Column": 28, "Meaning": 110}, freeze="A2")

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


# ---------------------------------------------------------------- CLI

def parse_sim_deck(text):
    """Parse an OPTCGSim export: one card per line, like '4xOP12-015'."""
    counts = Counter()
    for line in text.splitlines():
        m = re.match(r"\s*(\d+)\s*x\s*([A-Z0-9]+-\d+)", line, re.I)
        if m:
            counts[m.group(2).upper()] += int(m.group(1))
    return counts


def cmd_sync(args):
    db = connect()
    load_cards(db, args.refresh_cards)
    sync(db, args)


def cmd_analyze(args):
    db = connect()
    cards = load_cards(db, args.refresh_cards)
    if not args.no_sync:
        sync(db, args)
    entries, eras = load_entries(db, args)
    if not entries:
        sys.exit("No decklists matched your filters. Try a different --set, --event or --top.")
    parts = [f"events matching {' / '.join(args.event)}", f"set {args.set}",
             f"top {args.top}" if args.top else "all players"]
    parts += [f"last {args.days} days"] if args.days else []
    parts += [f"leader {args.leader}"] if args.leader else []
    parts += [f"name contains {args.format}"] if args.format else []
    out = Path(args.out or f"reports/optcg_{datetime.now():%Y-%m-%d}.xlsx")
    write_workbook(entries, cards, out, eras, ", ".join(parts))
    print(f"Wrote {len(entries)} decks from {len({e['tournament_id'] for e in entries})} tournaments "
          f"across {len({e['set'] for e in entries})} set(s) to {out}")


def cmd_deck(args):
    db = connect()
    cards = load_cards(db, args.refresh_cards)
    counts = parse_sim_deck(Path(args.file).read_text())
    leader = next((cid for cid in counts if cards.get(cid, {}).get("type") == "Leader"), None)
    if leader:
        counts.pop(leader)
    name = f"{cards[leader]['name']} ({leader})" if leader else "Unknown leader"
    entry = {"set": "", "leader_name": name, "player": "(you)", "placing": None, "record": "",
             "tournament": Path(args.file).name, "date": f"{datetime.now():%Y-%m-%d}", "counts": counts}
    out = Path(args.out or f"reports/{Path(args.file).stem}.xlsx")
    write_workbook([entry], cards, out)
    s, curve = entry["stats"], entry["curve"]
    print(f"{name}: {s['total']} cards | bricks {s['bricks']} | counters {s['with_counter']} "
          f"(+1k {s['counter_1k']}, +2k {s['counter_2k']}) | searchers {s['searchers']} | events {s['events']}")
    print("Cost curve: " + "  ".join(f"{i}:{curve.get(i, 0)}" for i in range(MAX_COST + 1)))
    if entry["unknown"]:
        print(f"Unknown cards (not counted): {', '.join(entry['unknown'])}")
    print(f"Wrote {out}")


def cmd_sql(args):
    db = connect()
    try:
        cur = db.execute(args.query)
    except sqlite3.Error as e:
        sys.exit(f"SQL error: {e}")
    rows = cur.fetchmany(args.limit)
    if not cur.description:
        db.commit()
        return
    cols = [d[0] for d in cur.description]
    table = [cols] + [["" if v is None else str(v) for v in r] for r in rows]
    widths = [min(max(len(row[i]) for row in table), 40) for i in range(len(cols))]
    for i, row in enumerate(table):
        print("  ".join(v[:40].ljust(w) for v, w in zip(row, widths)))
        if i == 0:
            print("  ".join("-" * w for w in widths))


def _filters(p):
    p.add_argument("--event", action="append", help="tournament name contains this (repeatable, default: chinoize)")
    p.add_argument("--format", help="only tournaments whose name contains this, e.g. EGB")
    p.add_argument("--days", type=int, help="only the last N days (default: no limit)")
    p.add_argument("--min-players", type=int, default=16, help="skip smaller tournaments (default 16)")


def main():
    ap = argparse.ArgumentParser(prog="optcg", description="Analyze One Piece TCG tournament decks into an Excel report.")
    ap.add_argument("--refresh-cards", action="store_true", help="re-download the card database (auto-refreshes weekly)")
    sub = ap.add_subparsers(dest="cmd", required=True)

    a = sub.add_parser("analyze", help="sync, then write an Excel report from the database")
    _filters(a)
    a.add_argument("--set", default="all", help="all (default), current, or a set code like OP16")
    a.add_argument("--top", type=int, default=8, help="only decks that placed this or better (default 8; 0 = everyone)")
    a.add_argument("--leader", help="only include this leader (name or card ID, partial match)")
    a.add_argument("--out", help="output .xlsx path (default reports/optcg_<date>.xlsx)")
    a.add_argument("--no-sync", action="store_true", help="use only what's already in the database (works offline)")
    a.set_defaults(func=cmd_analyze)

    s = sub.add_parser("sync", help="download new tournaments into the database without writing a report")
    _filters(s)
    s.set_defaults(func=cmd_sync)

    d = sub.add_parser("deck", help="analyze one decklist exported from OPTCGSim")
    d.add_argument("file", help="text file with lines like 4xOP12-015")
    d.add_argument("--out", help="output .xlsx path (default reports/<file>.xlsx)")
    d.set_defaults(func=cmd_deck)

    q = sub.add_parser("sql", help="run a SQL query against the database, e.g. on the deck_summary view")
    q.add_argument("query")
    q.add_argument("--limit", type=int, default=50, help="max rows to print (default 50)")
    q.set_defaults(func=cmd_sql)

    args = ap.parse_args()
    if args.cmd in ("analyze", "sync") and not args.event:
        args.event = ["chinoize"]
    args.func(args)


if __name__ == "__main__":
    main()
